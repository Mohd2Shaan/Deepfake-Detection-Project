"""Build the presentation report/Deepfake_Detection_Presentation.pptx from report/figures/.

    python report/make_figures.py
    python report/build_slides.py
"""

import re
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
FIG = HERE / "figures"
OUT = HERE / "Deepfake_Detection_Presentation.pptx"

ACCENT = RGBColor(0x1F, 0x5F, 0xAE)
ACCENT_SOFT = RGBColor(0xEA, 0xF2, 0xFB)
INK = RGBColor(0x1F, 0x23, 0x28)
MUTED = RGBColor(0x59, 0x63, 0x6E)
RULE = RGBColor(0xD8, 0xDE, 0xE4)
ORANGE = RGBColor(0xEB, 0x68, 0x34)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Segoe UI"
NUMERIC = re.compile(r"^[≈~]?\s*[\d.,]+\s*(%|M|k)?$")   # numbers are centred in tables, text is left-aligned

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]
slide_no = 0


# ---------------------------------------------------------------- helpers
def rect(slide, x, y, w, h, fill, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
    shp.shadow.inherit = False
    return shp


def text(slide, x, y, w, h, content, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """content: str or list of paragraphs; a paragraph is str or list of (text, bold) runs."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.03)
    paras = content if isinstance(content, list) else [content]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        runs = para if isinstance(para, list) else [(para, bold)]
        for t, b in runs:
            r = p.add_run(); r.text = t
            r.font.size = Pt(size); r.font.bold = b; r.font.color.rgb = color; r.font.name = FONT
    return tb


def bullets(slide, x, y, w, h, items, size=16, gap=6):
    """items: list of str or (bold_lead, rest)."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        lead, rest = (it if isinstance(it, tuple) else ("", it))
        r = p.add_run(); r.text = "•  "; r.font.size = Pt(size); r.font.color.rgb = ACCENT; r.font.name = FONT; r.font.bold = True
        if lead:
            r = p.add_run(); r.text = lead + " "; r.font.bold = True
            r.font.size = Pt(size); r.font.color.rgb = INK; r.font.name = FONT
        r = p.add_run(); r.text = rest; r.font.size = Pt(size); r.font.color.rgb = INK; r.font.name = FONT
    return tb


def image(slide, name, x, y, w, h, caption=None):
    """Place an image centred inside the box (x, y, w, h), keeping its aspect ratio."""
    path = FIG / name
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    pw, ph = int(iw * scale), int(ih * scale)
    px, py = x + (w - pw) // 2, y + (h - ph) // 2
    slide.shapes.add_picture(str(path), px, py, pw, ph)
    if caption:
        text(slide, x, py + ph + Inches(0.02), w, Inches(0.35), caption, size=11, color=MUTED, align=PP_ALIGN.CENTER)


def table(slide, x, y, w, rows, col_w=None, size=13, highlight=None, row_h=0.42):
    nr, nc = len(rows), len(rows[0])
    shp = slide.shapes.add_table(nr, nc, x, y, w, Inches(row_h * nr))
    tbl = shp.table
    if col_w:
        total = sum(col_w)
        for j, cw in enumerate(col_w):
            tbl.columns[j].width = Emu(int(w * cw / total))
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.fill.solid()
            cell.fill.fore_color.rgb = ACCENT if i == 0 else (ACCENT_SOFT if highlight == i else WHITE)
            tf = cell.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; p.text = ""
            r = p.add_run(); r.text = str(val); r.font.name = FONT; r.font.size = Pt(size)
            r.font.bold = i == 0 or highlight == i
            r.font.color.rgb = WHITE if i == 0 else INK
            p.alignment = PP_ALIGN.CENTER if (j > 0 and NUMERIC.match(str(val))) else PP_ALIGN.LEFT
    return tbl


def callout(slide, x, y, w, h, content, size=15, fill=ACCENT_SOFT, bar=ACCENT):
    rect(slide, x, y, w, h, fill)
    rect(slide, x, y, Inches(0.08), h, bar)
    text(slide, x + Inches(0.2), y, w - Inches(0.3), h, content, size=size, anchor=MSO_ANCHOR.MIDDLE)


def new_slide(title, kicker=""):
    global slide_no
    slide_no += 1
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, Inches(0.18), SH, ACCENT)
    if kicker:
        text(s, Inches(0.6), Inches(0.32), Inches(10), Inches(0.35), kicker.upper(), size=12, color=ACCENT, bold=True)
    text(s, Inches(0.6), Inches(0.62), Inches(12), Inches(0.75), title, size=30, bold=True)
    rect(s, Inches(0.62), Inches(1.38), Inches(1.2), Inches(0.05), ACCENT)
    text(s, Inches(0.6), SH - Inches(0.5), Inches(8), Inches(0.35), "Deepfake Detection · Real vs StyleGAN Faces",
         size=10, color=MUTED)
    text(s, SW - Inches(1.4), SH - Inches(0.5), Inches(0.9), Inches(0.35), str(slide_no), size=10, color=MUTED,
         align=PP_ALIGN.RIGHT)
    return s


L, T = Inches(0.6), Inches(1.65)          # content left / top
CW = SW - Inches(1.2)                      # content width
CH = SH - T - Inches(0.7)                  # content height

# ================================================================ 1 title
slide_no += 1
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, Inches(5.0), SH, ACCENT)
text(s, Inches(0.6), Inches(0.7), Inches(4), Inches(0.4), "MACHINE LEARNING · SEMESTER 5", size=13, color=WHITE, bold=True)
text(s, Inches(0.6), Inches(1.4), Inches(4.1), Inches(3), ["Deepfake", "Detection"], size=48, color=WHITE, bold=True)
text(s, Inches(0.6), Inches(3.4), Inches(4.1), Inches(1.2), "Real vs StyleGAN-generated faces with transfer learning and Grad-CAM",
     size=17, color=WHITE)
text(s, Inches(0.6), Inches(5.6), Inches(4.1), Inches(1.4),
     [[("Mohd Shaan", True)], "Submitted to Dr. Prakash", "October 2026"], size=15, color=WHITE)
image(s, "gradcam_fake.jpg", Inches(5.5), Inches(0.5), Inches(3.6), Inches(6.5))
tiles = [("99.84%", "test accuracy"), ("1.0000", "ROC-AUC"), ("33", "errors / 20,000"), ("140k", "face images")]
for i, (v, lbl) in enumerate(tiles):
    y = Inches(0.75 + i * 1.55)
    rect(s, Inches(9.6), y, Inches(3.2), Inches(1.3), ACCENT_SOFT)
    text(s, Inches(9.6), y + Inches(0.12), Inches(3.2), Inches(0.7), v, size=32, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)
    text(s, Inches(9.6), y + Inches(0.82), Inches(3.2), Inches(0.4), lbl, size=13, color=MUTED, align=PP_ALIGN.CENTER)

# ================================================================ 2 problem
s = new_slide("Problem and motivation", "Introduction")
bullets(s, L, T, Inches(6.2), CH, [
    ("Deepfakes:", "images or videos generated by deep networks that look authentic."),
    ("Entire-face synthesis:", "StyleGAN creates faces of people who do not exist — used for fake profiles, fraud and disinformation."),
    ("Humans can't tell:", "real and generated faces look the same to the eye."),
    ("Our task:", "binary classification — given one face image, predict REAL or FAKE with a probability P(fake)."),
], size=17, gap=12)
callout(s, Inches(7.2), T, Inches(5.5), Inches(4.6), [
    [("Objectives", True)], "",
    "1. Clean, verified data pipeline + EDA",
    "2. Fine-tune EfficientNet-B0 and ResNet-18",
    "3. Rigorous test-set evaluation + error analysis",
    "4. Explain decisions with Grad-CAM",
    "5. Streamlit demo + out-of-distribution test",
], size=17)

# ================================================================ 3 dataset
s = new_slide("Dataset: 140k Real and Fake Faces", "Data")
table(s, L, T, Inches(6.0), [
    ["Split", "Real", "Fake", "Total"],
    ["train", "50,000", "50,000", "100,000"],
    ["valid", "10,000", "10,000", "20,000"],
    ["test", "10,000", "10,000", "20,000"],
], col_w=[1.2, 1, 1, 1], size=15)
bullets(s, L, T + Inches(2.1), Inches(6.0), Inches(3), [
    ("Real:", "70k FFHQ photographs (Flickr)"),
    ("Fake:", "70k StyleGAN images"),
    ("256×256, face-aligned, perfectly balanced"),
    ("Predefined splits used as-is →", "no data leakage; test set used once"),
], size=16)
image(s, "eda_samples_real.jpg", Inches(6.9), T - Inches(0.05), Inches(2.9), Inches(4.6), "real (FFHQ)")
image(s, "eda_samples_fake.jpg", Inches(9.95), T - Inches(0.05), Inches(2.9), Inches(4.6), "fake (StyleGAN)")

# ================================================================ 4 EDA
s = new_slide("Exploratory data analysis — clean data, no shortcuts", "Data")
image(s, "eda_file_size.png", L, T, Inches(5.6), Inches(2.9), "JPEG file size per class — almost full overlap")
image(s, "eda_mean_face_spectrum.png", Inches(6.4), T, Inches(3.0), Inches(4.8), "mean face and mean Fourier spectrum")
bullets(s, Inches(9.6), T, Inches(3.3), CH, [
    ("0", "corrupted files"), ("0", "exact duplicates (MD5)"), ("50/50", "balance in every split"),
    ("All", "256×256 RGB"), ("No compression shortcut", "(Δ ≈ 0.5 KB)"),
    ("Spectra nearly identical", "→ signal is subtle and local → CNN"),
], size=15, gap=8)
bullets(s, L, T + Inches(3.45), Inches(5.6), Inches(1.6), [
    "Full scan of all 140,000 images on Kaggle (≈17 min)",
    "Mean fake face is sharper → StyleGAN faces more uniform in pose",
], size=14)

# ================================================================ 5 pipeline
s = new_slide("Pipeline and implementation", "Methodology")
image(s, "pipeline.png", L, T, CW, Inches(3.4))
table(s, L, T + Inches(3.6), CW, [
    ["Module", "dataset.py", "eda.py", "model.py", "train.py", "evaluate.py", "gradcam.py", "app.py"],
    ["Role", "load, resize, augment", "data checks", "backbones + head", "2-phase training", "metrics + errors", "heatmaps", "Streamlit demo"],
], size=12, row_h=0.45)
text(s, L, T + Inches(4.75), CW, Inches(0.4),
     "Built module by module · one commit + one theory document per module · trained on Kaggle (Tesla T4) · reproducible notebooks",
     size=13, color=MUTED)

# ================================================================ 6 preprocessing
s = new_slide("Preprocessing and augmentation", "Methodology")
table(s, L, T, Inches(7.2), [
    ["Step", "What", "Why"],
    ["Resize", "224 × 224", "pretraining resolution; 23% cheaper than 256²"],
    ["Normalise", "ImageNet mean / std", "match the input distribution of pretrained filters"],
    ["Flip", "horizontal, p = 0.5", "a mirrored face is still real/fake"],
    ["Rotate", "± 10°", "natural head tilt"],
    ["Colour jitter", "brightness / contrast ± 20%", "lighting variety, no colour shortcut"],
    ["Valid / test", "resize + normalise only", "deterministic, true distribution"],
], col_w=[1.2, 1.7, 3.2], size=13, row_h=0.5)
callout(s, L, T + Inches(3.9), Inches(7.2), Inches(1.1),
        "No blur or JPEG augmentation — it could erase the high-frequency GAN artifacts we want the model to learn.", size=14)
image(s, "eda_batch.jpg", Inches(8.1), T, Inches(4.6), Inches(4.9), "training batch after augmentation")

# ================================================================ 7 architectures
s = new_slide("Two ImageNet-pretrained CNNs", "Methodology")
table(s, L, T, CW, [
    ["", "EfficientNet-B0 (primary)", "ResNet-18 (comparison)"],
    ["Key idea", "compound scaling; MBConv + squeeze-and-excitation", "residual connections  y = F(x) + x"],
    ["Parameters", "4.0 M", "11.2 M"],
    ["Compute", "≈ 0.39 GFLOPs / image", "≈ 1.8 GFLOPs / image"],
    ["New head", "Linear(1280 → 1)", "Linear(512 → 1)"],
    ["Grad-CAM layer", "features[-1]  (1280 × 7 × 7)", "layer4[-1]  (512 × 7 × 7)"],
], col_w=[1.3, 2.6, 2.6], size=15, row_h=0.55)
bullets(s, L, T + Inches(3.7), CW, Inches(1.5), [
    ("Transfer learning:", "early layers already detect edges and textures — only adapt the later layers (fast, accurate, less over-fitting)."),
    ("Single logit:", "P(fake) = σ(z) — equivalent to a 2-class softmax, trained with BCEWithLogitsLoss."),
], size=15)

# ================================================================ 8 two-phase
s = new_slide("Two-phase transfer learning", "Training")
image(s, "two_phase.png", L, T - Inches(0.1), CW, Inches(1.9))
bullets(s, L, T + Inches(1.95), Inches(5.8), Inches(3), [
    ("Phase 1 — warm-up:", "backbone frozen, only the random new head learns (lr 1e-4)."),
    ("Phase 2 — fine-tune:", "unfreeze the last blocks with a new optimiser; small lr avoids destroying pretrained features."),
    ("Frozen BatchNorm", "kept in eval mode so its ImageNet statistics stay fixed."),
], size=15, gap=10)
table(s, Inches(6.7), T + Inches(2.0), Inches(6.0), [
    ["Configuration", "Phase-2 trainable", "Share"],
    ["EfficientNet, 2 blocks (run 02)", "1.13 M", "28%"],
    ["EfficientNet, 4 blocks (run 02b)", "3.70 M", "92%"],
    ["ResNet-18, layer3 + layer4", "10.49 M", "94%"],
], col_w=[3, 1.6, 0.9], size=14, highlight=2)

# ================================================================ 9 training setup
s = new_slide("Loss, optimiser and checkpointing", "Training")
cards = [
    ("BCE with logits", "BCE = −[ y·log p + (1−y)·log(1−p) ]\nconfident mistakes cost the most; sigmoid fused into the loss for stability"),
    ("Adam", "momentum + per-parameter adaptive step size\nlr 1e-4 (head) → 1e-5 / 1e-4 (fine-tune)"),
    ("ReduceLROnPlateau", "halve the lr when validation loss stops improving for 1 epoch"),
    ("Best checkpoint", "save weights whenever validation loss hits a new minimum — early stopping; resumable _last.pth every epoch"),
]
for i, (h, b) in enumerate(cards):
    x = L + (i % 2) * Inches(6.15); y = T + (i // 2) * Inches(2.45)
    rect(s, x, y, Inches(5.95), Inches(2.2), ACCENT_SOFT)
    text(s, x + Inches(0.25), y + Inches(0.15), Inches(5.5), Inches(0.5), h, size=19, bold=True, color=ACCENT)
    text(s, x + Inches(0.25), y + Inches(0.75), Inches(5.5), Inches(1.4), b.split("\n"), size=17)

# ================================================================ 10 training results
s = new_slide("Training results — validation accuracy", "Results")
image(s, "val_acc_runs.png", L, T, Inches(7.6), Inches(4.9))
bullets(s, Inches(8.5), T + Inches(0.1), Inches(4.3), CH, [
    ("Frozen phase plateaus < 80%", "— ImageNet features alone are not enough."),
    ("Fine-tuning jumps immediately", "(78.6% → 99.2% for EfficientNet at epoch 4)."),
    ("No over-fitting:", "train and validation stay close."),
    ("Best:", "EfficientNet-B0 99.88% · ResNet-18 98.96%."),
    ("≈ 7 min / epoch", "on a Kaggle Tesla T4."),
], size=15, gap=10)

# ================================================================ 11 ablation
s = new_slide("Ablation: how much of the network to fine-tune", "Results")
image(s, "curves_efficientnet_2blocks.png", L, T, Inches(6.0), Inches(2.6), "run 02 · 2 blocks, lr 1e-5 → 93.71% (under-fitting)")
image(s, "curves_efficientnet_4blocks.png", Inches(6.75), T, Inches(6.0), Inches(2.6), "run 02b · 4 blocks, lr 1e-4 → 99.88%")
callout(s, L, T + Inches(3.25), CW, Inches(1.6), [
    [("Lesson: ", True), ("the share of the network that is fine-tuned mattered more than the architecture. "
                         "With comparable trainable share (92% vs 94%), EfficientNet-B0 beats ResNet-18 with ~3× fewer parameters.", False)],
], size=16)

# ================================================================ 12 test results
s = new_slide("Test-set results (20,000 unseen images)", "Results")
table(s, L, T, CW, [
    ["Model", "Accuracy", "Precision", "Recall", "F1", "ROC-AUC", "Errors"],
    ["EfficientNet-B0 (final)", "99.84%", "99.79%", "99.88%", "99.84%", "1.0000", "33"],
    ["ResNet-18", "98.98%", "98.58%", "99.38%", "98.98%", "0.9994", "205"],
], col_w=[2.6, 1, 1, 1, 1, 1, 0.8], size=15, highlight=1, row_h=0.5)
image(s, "confusion_efficientnet.png", L, T + Inches(1.75), Inches(4.1), Inches(3.2), "EfficientNet-B0")
image(s, "confusion_resnet.png", Inches(4.75), T + Inches(1.75), Inches(4.1), Inches(3.2), "ResNet-18")
bullets(s, Inches(9.1), T + Inches(1.85), Inches(3.7), Inches(3.2), [
    ("Test ≈ validation", "→ no over-fitting"),
    ("6× fewer errors", "for EfficientNet"),
    ("More false positives", "than false negatives (21 vs 12)"),
], size=15, gap=10)

# ================================================================ 13 ROC + errors
s = new_slide("ROC curve and error count", "Results")
image(s, "roc_curve.png", L, T, Inches(5.2), Inches(4.9))
image(s, "test_errors.png", Inches(6.1), T + Inches(0.2), Inches(6.7), Inches(2.4))
callout(s, Inches(6.1), T + Inches(3.0), Inches(6.7), Inches(1.7),
        "AUC ≈ 1.000: almost every fake gets a higher P(fake) than almost every real photo — the result does not depend on the 0.5 threshold.",
        size=15)

# ================================================================ 14 error analysis
s = new_slide("Error analysis — the most confident mistakes", "Results")
image(s, "errors_efficientnet.jpg", L, T - Inches(0.1), Inches(5.3), Inches(5.1))
bullets(s, Inches(6.2), T, Inches(6.6), CH, [
    ("Real → fake (false positives):", "studio-like portraits — plain background, smooth or retouched skin, make-up, flash, symmetric frontal pose. Real photos that look as “clean” as StyleGAN."),
    ("Fake → real (false negatives):", "busy realistic backgrounds or clothing, unusual lighting, sunglasses, face paint — StyleGAN artifacts less visible."),
    ("Most errors have moderate confidence", "(65–90%); only a few exceed 95%."),
], size=16, gap=14)

# ================================================================ 15 grad-cam method
s = new_slide("Explainability: how Grad-CAM works", "Explainability")
bullets(s, L, T, Inches(6.4), CH, [
    ("1. Forward pass:", "keep the feature maps Aᵏ of the last conv layer (7 × 7)."),
    ("2. Backward pass:", "gradients of the score y w.r.t. Aᵏ  (y = z for fake, −z for real)."),
    ("3. Weights:", "αₖ = global average of the gradients."),
    ("4. Heatmap:", "ReLU( Σₖ αₖ Aᵏ ), upsample to 224 × 224, overlay."),
    ("Why the last conv layer:", "most class-specific features, still spatial."),
], size=16, gap=11)
image(s, "gradcam_manual_vs_library.png", Inches(7.2), T, Inches(5.6), Inches(2.4))
callout(s, Inches(7.2), T + Inches(2.7), Inches(5.6), Inches(1.5),
        "Our hand-written Grad-CAM matches the pytorch-grad-cam library: correlation 0.99999999.", size=15)

# ================================================================ 16 grad-cam findings
s = new_slide("Grad-CAM findings — where the model looks", "Explainability")
image(s, "gradcam_fake.jpg", L, T - Inches(0.1), Inches(2.9), Inches(5.2), "correct fakes")
image(s, "gradcam_real.jpg", Inches(3.65), T - Inches(0.1), Inches(2.9), Inches(5.2), "correct reals")
bullets(s, Inches(6.9), T, Inches(5.9), CH, [
    ("Fakes:", "tight, central attention — between the eyes, nose bridge, mouth/teeth, where generator artifacts concentrate."),
    ("Real photos:", "spread over hair, background, hats, hands and objects — natural camera detail StyleGAN renders poorly."),
    ("Errors follow the same logic:", "real→fake errors get face-centred attention; fake→real errors look at realistic backgrounds."),
    ("No trivial shortcut", "(border, watermark) — but the cues are specific to StyleGAN’s aligned faces."),
], size=15, gap=11)

# ================================================================ 17 demo + OOD
s = new_slide("Demo app and out-of-distribution test", "Deployment")
steps = ["Upload face image", "Resize + normalise", "EfficientNet-B0", "REAL / FAKE + confidence", "Grad-CAM overlay"]
for i, st in enumerate(steps):
    x = L + i * Inches(2.45)
    rect(s, x, T, Inches(2.2), Inches(0.95), ACCENT_SOFT)
    text(s, x, T, Inches(2.2), Inches(0.95), st, size=14, bold=True, color=ACCENT, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
bullets(s, L, T + Inches(1.25), Inches(6.0), Inches(3), [
    ("streamlit run app.py", "— plain interface: upload → predict → show result."),
    ("≈ 1 s per image", "including Grad-CAM, on a laptop CPU."),
    ("Model cached", "with @st.cache_resource."),
], size=15, gap=10)
callout(s, Inches(6.9), T + Inches(1.25), Inches(5.9), Inches(3.6), [
    [("Out-of-distribution test", True)], "",
    "3 face images generated with ChatGPT",
    [("→ all predicted REAL with ≈100% confidence", True)], "",
    "The model learned StyleGAN’s fingerprint, not “fakeness” in general.",
], size=16, fill=RGBColor(0xFF, 0xF6, 0xF1), bar=ORANGE)

# ================================================================ 18 discussion
s = new_slide("Discussion", "Discussion")
callout(s, L, T, Inches(5.9), Inches(4.9), [
    [("Is 99.8% over-fitting?  No.", True)], "",
    "Train 99.72% · validation 99.88% · test 99.84% — within 0.2% on 40,000 images never trained on.",
    "",
    "Accuracy is high because the task is narrow: one generator (StyleGAN) with a consistent fingerprint, all faces aligned.",
], size=15)
callout(s, Inches(6.85), T, Inches(5.9), Inches(4.9), [
    [("Why does it fail on ChatGPT images?", True)], "",
    "1. Different generator → different fingerprint",
    "2. Only two classes → “real” means “not StyleGAN”",
    "3. Larger, un-aligned images → different input statistics",
    "4. Networks are over-confident outside their training data",
], size=15, fill=RGBColor(0xFF, 0xF6, 0xF1), bar=ORANGE)

# ================================================================ 19 limitations
s = new_slide("Limitations", "Discussion")
table(s, L, T, CW, [
    ["Limitation", "Consequence"],
    ["Single generator (StyleGAN) in training", "fails on other generators — shown with ChatGPT images"],
    ["Aligned 256 px FFHQ-style faces only", "arbitrary photos are out of distribution; no face detector in the app"],
    ["No robustness tests", "JPEG re-compression, resizing, blur, screenshots may erase the fingerprint"],
    ["Uncalibrated confidence", "high confidence on unfamiliar inputs; no “unknown” option"],
    ["Coarse 7 × 7 Grad-CAM", "shows roughly where, not which artifact"],
    ["Still images only", "no temporal cues for video deepfakes / face swaps"],
], col_w=[2.2, 3.8], size=15, row_h=0.6)

# ================================================================ 20 future work
s = new_slide("Future work and extensions", "Next steps")
cards = [
    ("Multi-generator data", "add StyleGAN2/3 + diffusion fakes (Stable Diffusion, Midjourney, DALL·E); leave-one-generator-out evaluation"),
    ("General features", "frozen CLIP/ViT features + linear probe (Ojha et al., 2023)"),
    ("Frequency branch", "two-stream RGB + Fourier model for up-sampling artifacts"),
    ("Robustness", "train with JPEG, blur, resize, noise augmentation"),
    ("Face detection + calibration", "MTCNN/RetinaFace in the app; temperature scaling and an “uncertain” output"),
    ("Video and scale", "FaceForensics++/DFDC with temporal models; ONNX + REST API deployment"),
]
for i, (h, b) in enumerate(cards):
    x = L + (i % 3) * Inches(4.1); y = T + (i // 3) * Inches(2.45)
    rect(s, x, y, Inches(3.9), Inches(2.2), ACCENT_SOFT)
    text(s, x + Inches(0.2), y + Inches(0.15), Inches(3.5), Inches(0.5), h, size=17, bold=True, color=ACCENT)
    text(s, x + Inches(0.2), y + Inches(0.7), Inches(3.5), Inches(1.45), b, size=16)

# ================================================================ 21 conclusion
s = new_slide("Conclusion", "Summary")
bullets(s, L, T, Inches(7.2), CH, [
    ("Complete pipeline:", "verified data → two-phase transfer learning → rigorous evaluation → Grad-CAM → demo app."),
    ("EfficientNet-B0:", "99.84% test accuracy, ROC-AUC 1.0000, 33 errors in 20,000 — 6× fewer than ResNet-18 with 3× fewer parameters."),
    ("Fine-tuning depth matters:", "28% → 92% of weights trained took EfficientNet from 93.7% to 99.9%."),
    ("Interpretable:", "inner-face artifacts for fakes, natural context for real photos."),
    ("Key lesson:", "near-perfect accuracy on one generator ≠ a general deepfake detector."),
], size=16, gap=12)
rect(s, Inches(8.3), T, Inches(4.45), CH, ACCENT)
text(s, Inches(8.3), T + Inches(1.2), Inches(4.45), Inches(1), "Thank you", size=40, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
text(s, Inches(8.3), T + Inches(2.3), Inches(4.45), Inches(0.6), "Questions?", size=22, color=WHITE, align=PP_ALIGN.CENTER)
text(s, Inches(8.3), T + Inches(3.6), Inches(4.45), Inches(0.8),
     "github.com/Mohd2Shaan/\nDeepfake-Detection-Project".split("\n"), size=13, color=WHITE, align=PP_ALIGN.CENTER)

prs.save(OUT)
print("written", OUT, f"({OUT.stat().st_size / 1e6:.1f} MB, {slide_no} slides)")
